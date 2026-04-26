"""Format-specific readers. All parsers return a list of dicts with at minimum
{text, rating, date, source, language?, branch?, reviewer?}.
"""
from __future__ import annotations

import csv
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

# Optional dependency: openpyxl for .xlsx
try:
    from openpyxl import load_workbook  # type: ignore
    HAS_XLSX = True
except ImportError:
    HAS_XLSX = False


RATING_KEYS = ["rating", "stars", "star_rating", "starrating", "score", "review_rating", "overall_rating"]
TEXT_KEYS = ["text", "review", "review_text", "comment", "comments", "feedback", "body", "message", "content"]
DATE_KEYS = ["date", "review_date", "created_at", "createdat", "timestamp", "time", "submitted_at",
             "create_time", "createtime", "update_time", "updatetime", "publishedat", "published_at"]
BRANCH_KEYS = ["branch", "outlet", "location", "store", "restaurant", "site", "venue"]
REVIEWER_KEYS = ["reviewer", "name", "user", "customer", "author", "reviewer_name", "displayname"]

# Google review exports use word-form ratings instead of digits.
_WORD_RATINGS = {
    "ZERO": 0, "ONE": 1, "TWO": 2, "THREE": 3, "FOUR": 4, "FIVE": 5,
    "STAR_RATING_UNSPECIFIED": None,
}

# Google reviews often append a translated copy. We prefer the original.
_TRANSLATED_RE = re.compile(
    r"\(Translated by Google\)(?P<translated>.*?)\(Original\)(?P<original>.*)$",
    re.DOTALL | re.IGNORECASE,
)


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", s.lower())


def _pick(row: dict, candidates: list[str]) -> Any:
    norm_map = {_norm(k): k for k in row.keys()}
    for c in candidates:
        if _norm(c) in norm_map:
            return row[norm_map[_norm(c)]]
    return None


def _coerce_rating(value: Any, scale_hint: int = 5) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    # Google review exports use word-form ratings: "FIVE", "FOUR", etc.
    upper = s.upper()
    if upper in _WORD_RATINGS:
        return None if _WORD_RATINGS[upper] is None else float(_WORD_RATINGS[upper])
    m = re.search(r"(\d+(?:\.\d+)?)", s)
    if not m:
        return None
    n = float(m.group(1))
    # Sometimes ratings are "4 out of 5" or "8/10"
    if "/" in s:
        parts = s.split("/")
        try:
            denom = float(re.search(r"(\d+(?:\.\d+)?)", parts[1]).group(1))
            if denom and denom != scale_hint:
                # Normalize to scale_hint
                n = (n / denom) * scale_hint
        except Exception:
            pass
    return n


def _flatten_reviewer(value: Any) -> str | None:
    """Google review exports often nest reviewer as {displayName: ..., profilePhotoUrl: ...}."""
    if value is None or value == "":
        return None
    if isinstance(value, dict):
        return value.get("displayName") or value.get("name") or None
    return str(value).strip()


def _split_translated_comment(text: str) -> str:
    """If the comment has '(Translated by Google) ... (Original) ...', keep the original.
    Otherwise return as-is."""
    if not text:
        return text
    m = _TRANSLATED_RE.search(text)
    if m:
        return m.group("original").strip()
    return text


def _coerce_date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    s = str(value).strip()
    # Normalize trailing 'Z' (UTC indicator) so plain ISO formats parse it.
    if s.endswith("Z"):
        s = s[:-1]
    fmts = [
        "%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%S.%f",
        "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M:%S.%f",
        "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y",
        "%b %d, %Y", "%B %d, %Y", "%d %b %Y", "%d %B %Y",
    ]
    for f in fmts:
        try:
            return datetime.strptime(s, f).date().isoformat()
        except ValueError:
            continue
    # Epoch seconds or ms
    if s.isdigit():
        n = int(s)
        try:
            if n > 10**12:  # ms
                return datetime.utcfromtimestamp(n / 1000).date().isoformat()
            if n > 10**9:
                return datetime.utcfromtimestamp(n).date().isoformat()
        except (OverflowError, OSError, ValueError):
            pass
    return None


def _row_to_record(row: dict, source: str, overrides: dict | None = None) -> dict | None:
    overrides = overrides or {}

    def col(key, fallback_keys):
        if overrides.get(key):
            ov = overrides[key]
            if ov in row:
                return row[ov]
            # Try case-insensitive
            for k in row:
                if _norm(k) == _norm(ov):
                    return row[k]
        return _pick(row, fallback_keys)

    text = col("text_column", TEXT_KEYS)
    rating = col("rating_column", RATING_KEYS)
    date = col("date_column", DATE_KEYS)
    branch = col("branch_column", BRANCH_KEYS)
    reviewer = col("reviewer_column", REVIEWER_KEYS)

    text = _split_translated_comment(str(text).strip() if text else "")
    rating_val = _coerce_rating(rating, scale_hint=overrides.get("rating_scale", 5))
    date_val = _coerce_date(date)

    if not text and rating_val is None:
        return None

    return {
        "text": text,
        "rating": rating_val,
        "date": date_val,
        "branch": (str(branch).strip() if branch else None),
        "reviewer": _flatten_reviewer(reviewer),
        "source": source,
    }


def parse_csv(path: Path, overrides: dict | None = None) -> list[dict]:
    rows = []
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for r in reader:
            rec = _row_to_record(r, source=path.name, overrides=overrides)
            if rec:
                rows.append(rec)
    return rows


def parse_xlsx(path: Path, overrides: dict | None = None) -> list[dict]:
    if not HAS_XLSX:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")
    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        headers = None
        for row in ws.iter_rows(values_only=True):
            if headers is None:
                headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(row)]
                continue
            if all(c is None for c in row):
                continue
            r = dict(zip(headers, row))
            rec = _row_to_record(r, source=f"{path.name}::{ws.title}", overrides=overrides)
            if rec:
                rows.append(rec)
    return rows


def parse_json(path: Path, overrides: dict | None = None) -> list[dict]:
    data = json.loads(path.read_text(encoding="utf-8"))
    rows = []

    # Google Takeout My Activity format detection
    if isinstance(data, list) and data and isinstance(data[0], dict) and "title" in data[0] and "time" in data[0]:
        for item in data:
            title = item.get("title", "")
            if "review" not in title.lower() and "rated" not in title.lower():
                continue
            text = ""
            details = item.get("details", [])
            for d in details if isinstance(details, list) else []:
                if isinstance(d, dict) and d.get("name"):
                    text = text + " " + str(d.get("name"))
            # Star rating sometimes embedded in title: "Rated 4 stars on ..."
            m = re.search(r"(\d)\s*stars?", title, re.IGNORECASE)
            rating = float(m.group(1)) if m else None
            rows.append({
                "text": text.strip() or item.get("snippet", "") or "",
                "rating": rating,
                "date": _coerce_date(item.get("time")),
                "branch": None,
                "reviewer": None,
                "source": "google-takeout",
            })
        return [r for r in rows if r["text"] or r["rating"] is not None]

    # Generic JSON list of objects
    if isinstance(data, list):
        for r in data:
            if isinstance(r, dict):
                rec = _row_to_record(r, source=path.name, overrides=overrides)
                if rec:
                    rows.append(rec)
        return rows

    # JSON object with a key holding the list
    if isinstance(data, dict):
        for k, v in data.items():
            if isinstance(v, list) and v and isinstance(v[0], dict):
                for r in v:
                    rec = _row_to_record(r, source=f"{path.name}::{k}", overrides=overrides)
                    if rec:
                        rows.append(rec)
        return rows

    return rows


def parse_path(path: Path, overrides: dict | None = None) -> list[dict]:
    if path.is_dir():
        out = []
        for child in sorted(path.iterdir()):
            if child.is_file() and child.suffix.lower() in {".csv", ".xlsx", ".json"}:
                out.extend(parse_path(child, overrides))
        return out
    suf = path.suffix.lower()
    if suf == ".csv":
        return parse_csv(path, overrides)
    if suf == ".xlsx":
        return parse_xlsx(path, overrides)
    if suf == ".json":
        return parse_json(path, overrides)
    raise ValueError(f"Unsupported file type: {path}")


def inspect(path: Path) -> dict:
    """Return detected columns and a sample row without running full pipeline."""
    suf = path.suffix.lower()
    if suf == ".csv":
        with open(path, "r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            sample = next(iter(reader), None)
            return {"columns": list(sample.keys()) if sample else [], "sample": sample}
    if suf == ".xlsx":
        if not HAS_XLSX:
            return {"error": "openpyxl not installed"}
        wb = load_workbook(filename=str(path), read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = next(rows, [])
        sample = next(rows, [])
        return {
            "columns": [str(h) for h in headers],
            "sample": dict(zip([str(h) for h in headers], sample)),
            "sheets": [s.title for s in wb.worksheets],
        }
    if suf == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, list) and data:
            return {"format": "list", "count": len(data), "sample": data[0]}
        if isinstance(data, dict):
            return {"format": "object", "keys": list(data.keys())}
    return {"error": "unknown format"}
